"""
Utilitaire pour parser des référentiels de formulaires depuis Excel
Format Excel supporté avec plusieurs feuilles:
- Metadata: informations du formulaire
- Sections: définition des sections
- Fields: définition des champs
- Options: options pour select/radio/checkbox
- Validation: règles de validation
"""

import openpyxl
from typing import Dict, Any, List, Optional
import json


class ExcelReferentielParser:
    """Parse un fichier Excel pour créer un référentiel de formulaire"""
    
    @staticmethod
    def parse_excel_file(file_path: str) -> Dict[str, Any]:
        """
        Parse un fichier Excel et retourne un référentiel de formulaire.
        Détecte automatiquement le format (multi-feuilles ou simple).
        
        Args:
            file_path: Chemin vers le fichier Excel
            
        Returns:
            Dictionnaire avec la structure du référentiel
        """
        wb = openpyxl.load_workbook(file_path, data_only=True)
        
        try:
            # Détecter le format du fichier
            sheet_names = wb.sheetnames
            
            # Format simple: une seule feuille avec colonnes Section, Groupe, Champ, Type, Label, etc.
            if len(sheet_names) == 1 or "Metadata" not in sheet_names:
                print("🔍 Format détecté: SIMPLE (une feuille)")
                return ExcelReferentielParser._parse_simple_format(wb)
            
            # Format complexe: plusieurs feuilles (Metadata, Sections, Fields, etc.)
            print("🔍 Format détecté: COMPLEXE (multi-feuilles)")
            return ExcelReferentielParser._parse_complex_format(wb)
            
        finally:
            # Fermer le workbook pour libérer le fichier
            wb.close()
    
    @staticmethod
    def _parse_simple_format(wb: openpyxl.Workbook) -> Dict[str, Any]:
        """
        Parse un fichier Excel au format simple (une seule feuille).
        Format attendu: Section | Groupe | Champ | Type | Label | Requis | Options
        """
        ws = wb.active
        
        # Lire les en-têtes
        headers = [cell.value for cell in ws[1]]
        print(f"📋 En-têtes trouvés: {headers}")
        
        # Trouver les indices des colonnes
        col_map = {}
        for i, header in enumerate(headers):
            if header:
                col_map[str(header).strip().lower()] = i
        
        print(f"🗺️  Mapping colonnes: {col_map}")
        
        # Parser les données
        sections_dict = {}  # section_title -> section_data
        
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            # Extraire les valeurs
            section_title = row[col_map.get('section', 0)] if 'section' in col_map else None
            groupe = row[col_map.get('groupe', 1)] if 'groupe' in col_map else None
            champ_name = row[col_map.get('champ', 2)] if 'champ' in col_map else None
            champ_type = row[col_map.get('type', 3)] if 'type' in col_map else None
            label = row[col_map.get('label', 4)] if 'label' in col_map else None
            requis = row[col_map.get('requis', 5)] if 'requis' in col_map else None
            options = row[col_map.get('options', 6)] if 'options' in col_map else None
            
            # Skip les lignes vides
            if not champ_name or not champ_type:
                continue
            
            print(f"📝 Ligne {row_idx}: Section='{section_title}', Champ='{champ_name}', Type='{champ_type}'")
            
            # Utiliser une section par défaut si pas spécifiée
            if not section_title:
                section_title = "Informations générales"
            
            # Créer la section si elle n'existe pas
            if section_title not in sections_dict:
                sections_dict[section_title] = {
                    "id": section_title.lower().replace(" ", "_"),
                    "title": section_title,
                    "description": "",
                    "groups": []
                }
            
            # Construire le champ
            is_required = str(requis).strip().lower() in ['oui', 'yes', 'true', '1'] if requis else False
            
            field = {
                "id": champ_name,
                "name": champ_name,
                "type": str(champ_type).strip().lower(),
                "label": label or champ_name,
                "required": is_required
            }
            
            # Ajouter les règles de validation sous forme de liste
            if is_required:
                validations = []
                
                # Règle required
                validations.append({
                    "type": "required",
                    "message": f"{label or champ_name} est obligatoire"
                })
                
                # Validations spécifiques par type
                if champ_type == "email":
                    validations.append({
                        "type": "pattern",
                        "pattern": r"^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$",
                        "message": "Adresse email invalide"
                    })
                elif champ_type == "tel":
                    validations.append({
                        "type": "pattern",
                        "pattern": r"^[0-9+\s\-()]{10,}$",
                        "message": "Numéro de téléphone invalide"
                    })
                elif champ_type == "text" and "nom" in champ_name.lower():
                    validations.append({
                        "type": "minLength",
                        "value": 2,
                        "message": "Minimum 2 caractères"
                    })
                elif champ_type == "number":
                    validations.append({
                        "type": "min",
                        "value": 0,
                        "message": "La valeur doit être positive"
                    })
                
                field["validation"] = validations
            
            # Ajouter les options si présentes
            if options and champ_type in ['select', 'radio', 'multiselect', 'checkbox-group']:
                # Split par | ou ,
                options_list = str(options).split('|') if '|' in str(options) else str(options).split(',')
                field["options"] = [{"label": opt.strip(), "value": opt.strip()} for opt in options_list if opt.strip()]
            
            # Ajouter au groupe ou directement à la section
            if groupe:
                # Trouver ou créer le groupe
                group = None
                for g in sections_dict[section_title]["groups"]:
                    if g["title"] == groupe:
                        group = g
                        break
                
                if not group:
                    group = {
                        "id": groupe.lower().replace(" ", "_").replace("'", ""),
                        "title": groupe,
                        "fields": []
                    }
                    sections_dict[section_title]["groups"].append(group)
                
                group["fields"].append(field)
            else:
                # Pas de groupe, créer un groupe par défaut
                if not sections_dict[section_title]["groups"]:
                    sections_dict[section_title]["groups"].append({
                        "id": "default",
                        "title": "",
                        "fields": []
                    })
                sections_dict[section_title]["groups"][0]["fields"].append(field)
        
        # Convertir en liste
        sections_list = list(sections_dict.values())
        
        print(f"\n✅ Parsing terminé:")
        print(f"   📊 {len(sections_list)} sections créées")
        total_fields = sum(len(g["fields"]) for s in sections_list for g in s["groups"])
        print(f"   📝 {total_fields} champs au total\n")
        
        # Construire la structure finale
        return {
            "version": "1.0.0",
            "metadata": {
                "name": "Formulaire importé",
                "description": "Formulaire importé depuis Excel (format simple)",
                "author": None,
                "createdAt": None,
                "tags": []
            },
            "config": {
                "id": "imported_form",
                "version": "1.0.0",
                "name": "Formulaire importé",
                "description": "Formulaire importé depuis Excel (format simple)",
                "sections": sections_list,
                "layout": {"type": "grid", "columns": 1},
                "submitButton": {"label": "Soumettre", "position": "center"},
                "cancelButton": {"label": "Annuler", "show": True},
                "validateOnBlur": True,
                "validateOnSubmit": True,
                "messages": {
                    "required": "Ce champ est obligatoire",
                    "invalid": "Valeur invalide",
                    "success": "Formulaire envoyé avec succès !",
                    "error": "Une erreur est survenue"
                }
            }
        }
    
    @staticmethod
    def _parse_complex_format(wb: openpyxl.Workbook) -> Dict[str, Any]:
        """Parse le format complexe avec plusieurs feuilles"""
        # Parser chaque feuille
        metadata = ExcelReferentielParser._parse_metadata(wb)
        sections = ExcelReferentielParser._parse_sections(wb)
        fields = ExcelReferentielParser._parse_fields(wb)
        options = ExcelReferentielParser._parse_options(wb)
        validations = ExcelReferentielParser._parse_validations(wb)
        
        # Construire le référentiel
        referentiel = {
        "version": metadata.get("version", "1.0.0"),
        "metadata": {
            "name": metadata.get("name", "Formulaire sans nom"),
            "description": metadata.get("description"),
            "author": metadata.get("author"),
            "createdAt": metadata.get("createdAt"),
            "tags": metadata.get("tags", "").split(",") if metadata.get("tags") else []
        },
        "config": {
            "id": metadata.get("id", "form"),
            "version": metadata.get("version", "1.0.0"),
            "name": metadata.get("name", "Formulaire sans nom"),
            "description": metadata.get("description"),
            "sections": ExcelReferentielParser._build_sections(sections, fields, options, validations),
            "layout": {
                "type": metadata.get("layout_type", "grid"),
                "columns": int(metadata.get("layout_columns", 1))
            },
            "submitButton": {
                "label": metadata.get("submit_label", "Soumettre"),
                "position": metadata.get("submit_position", "center")
            },
            "cancelButton": {
                "label": metadata.get("cancel_label", "Annuler"),
                "show": metadata.get("show_cancel", "true").lower() == "true"
            },
            "validateOnBlur": metadata.get("validate_on_blur", "true").lower() == "true",
            "validateOnSubmit": metadata.get("validate_on_submit", "true").lower() == "true",
            "messages": {
                "required": metadata.get("msg_required", "Ce champ est obligatoire"),
                "invalid": metadata.get("msg_invalid", "Valeur invalide"),
                "success": metadata.get("msg_success", "Formulaire envoyé avec succès !"),
                "error": metadata.get("msg_error", "Une erreur est survenue")
            }
        }
    }
        
        return referentiel
    
    @staticmethod
    def _parse_metadata(wb: openpyxl.Workbook) -> Dict[str, Any]:
        """Parse la feuille Metadata"""
        metadata = {}
        
        if "Metadata" in wb.sheetnames:
            ws = wb["Metadata"]
            
            # Format: Colonne A = clé, Colonne B = valeur
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] and row[1]:
                    metadata[str(row[0]).strip()] = row[1]
        
        return metadata
    
    @staticmethod
    def _parse_sections(wb: openpyxl.Workbook) -> List[Dict[str, Any]]:
        """Parse la feuille Sections"""
        sections = []
        
        if "Sections" in wb.sheetnames:
            ws = wb["Sections"]
            
            # Header: id | title | description | icon | order
            headers = [cell.value for cell in ws[1]]
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0]:  # Si ID existe
                    section = {}
                    for i, header in enumerate(headers):
                        if header and i < len(row) and row[i] is not None:
                            section[str(header).strip()] = row[i]
                    sections.append(section)
        
        return sections
    
    @staticmethod
    def _parse_fields(wb: openpyxl.Workbook) -> List[Dict[str, Any]]:
        """Parse la feuille Fields"""
        fields = []
        
        if "Fields" in wb.sheetnames:
            ws = wb["Fields"]
            
            # Header: id | name | type | label | placeholder | required | section_id | group_id | ...
            headers = [cell.value for cell in ws[1]]
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0]:  # Si ID existe
                    field = {}
                    for i, header in enumerate(headers):
                        if header and i < len(row) and row[i] is not None:
                            field[str(header).strip()] = row[i]
                    fields.append(field)
        
        return fields
    
    @staticmethod
    def _parse_options(wb: openpyxl.Workbook) -> Dict[str, List[Dict[str, Any]]]:
        """Parse la feuille Options"""
        options_by_field = {}
        
        if "Options" in wb.sheetnames:
            ws = wb["Options"]
            
            # Header: field_id | label | value | disabled
            headers = [cell.value for cell in ws[1]]
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0]:  # Si field_id existe
                    field_id = str(row[0]).strip()
                    option = {}
                    
                    for i, header in enumerate(headers):
                        if header and i < len(row) and row[i] is not None:
                            option[str(header).strip()] = row[i]
                    
                    if field_id not in options_by_field:
                        options_by_field[field_id] = []
                    options_by_field[field_id].append(option)
        
        return options_by_field
    
    @staticmethod
    def _parse_validations(wb: openpyxl.Workbook) -> Dict[str, List[Dict[str, Any]]]:
        """Parse la feuille Validation"""
        validations_by_field = {}
        
        if "Validation" in wb.sheetnames:
            ws = wb["Validation"]
            
            # Header: field_id | type | message | value | pattern
            headers = [cell.value for cell in ws[1]]
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0]:  # Si field_id existe
                    field_id = str(row[0]).strip()
                    validation = {}
                    
                    for i, header in enumerate(headers):
                        if header and i < len(row) and row[i] is not None:
                            validation[str(header).strip()] = row[i]
                    
                    if field_id not in validations_by_field:
                        validations_by_field[field_id] = []
                    validations_by_field[field_id].append(validation)
        
        return validations_by_field
    
    @staticmethod
    def _build_sections(
        sections: List[Dict[str, Any]],
        fields: List[Dict[str, Any]],
        options: Dict[str, List[Dict[str, Any]]],
        validations: Dict[str, List[Dict[str, Any]]]
    ) -> List[Dict[str, Any]]:
        """Construit la structure complète des sections avec champs"""
        
        built_sections = []
        
        for section in sections:
            section_id = section.get("id")
            
            # Grouper les champs par group_id
            groups_dict = {}
            
            for field in fields:
                if field.get("section_id") == section_id:
                    group_id = field.get("group_id", "default")
                    
                    if group_id not in groups_dict:
                        groups_dict[group_id] = {
                            "id": group_id,
                            "title": field.get("group_title"),
                            "layout": field.get("group_layout", "grid"),
                            "columns": int(field.get("group_columns", 1)),
                            "fields": []
                        }
                    
                    # Construire le champ
                    field_config = {
                        "id": field.get("id"),
                        "name": field.get("name"),
                        "type": field.get("type"),
                        "label": field.get("label"),
                        "placeholder": field.get("placeholder"),
                        "required": str(field.get("required", "false")).lower() == "true",
                        "disabled": str(field.get("disabled", "false")).lower() == "true",
                        "readonly": str(field.get("readonly", "false")).lower() == "true"
                    }
                    
                    # Ajouter les propriétés optionnelles
                    if field.get("description"):
                        field_config["description"] = field.get("description")
                    if field.get("tooltip"):
                        field_config["tooltip"] = field.get("tooltip")
                    if field.get("defaultValue"):
                        field_config["defaultValue"] = field.get("defaultValue")
                    if field.get("min"):
                        field_config["min"] = field.get("min")
                    if field.get("max"):
                        field_config["max"] = field.get("max")
                    if field.get("span") is not None:
                        field_config["span"] = int(field.get("span", 1))
                    
                    # Ajouter les options si disponibles
                    field_id = field.get("id")
                    if field_id in options:
                        field_config["options"] = [
                            {
                                "label": opt.get("label"),
                                "value": opt.get("value"),
                                "disabled": str(opt.get("disabled", "false")).lower() == "true"
                            }
                            for opt in options[field_id]
                        ]
                    
                    # Ajouter les validations si disponibles
                    if field_id in validations:
                        field_config["validation"] = [
                            {
                                "type": val.get("type"),
                                "message": val.get("message"),
                                "value": val.get("value"),
                                "pattern": val.get("pattern")
                            }
                            for val in validations[field_id]
                        ]
                    
                    groups_dict[group_id]["fields"].append(field_config)
            
            # Construire la section
            built_section = {
                "id": section_id,
                "title": section.get("title"),
                "description": section.get("description"),
                "icon": section.get("icon"),
                "order": int(section.get("order", 0)),
                "groups": list(groups_dict.values())
            }
            
            built_sections.append(built_section)
        
        return sorted(built_sections, key=lambda x: x.get("order", 0))


def create_excel_template(output_path: str):
    """
    Crée un template Excel vide avec toutes les feuilles nécessaires
    
    Args:
        output_path: Chemin où sauvegarder le template
    """
    wb = openpyxl.Workbook()
    
    # Supprimer la feuille par défaut
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
    
    # Feuille Metadata
    ws_meta = wb.create_sheet("Metadata")
    ws_meta.append(["Clé", "Valeur"])
    ws_meta.append(["id", "contact-form"])
    ws_meta.append(["version", "1.0.0"])
    ws_meta.append(["name", "Formulaire de Contact"])
    ws_meta.append(["description", "Un formulaire simple"])
    ws_meta.append(["author", ""])
    ws_meta.append(["tags", "contact,exemple"])
    ws_meta.append(["layout_type", "grid"])
    ws_meta.append(["layout_columns", "1"])
    ws_meta.append(["submit_label", "Envoyer"])
    ws_meta.append(["submit_position", "center"])
    ws_meta.append(["cancel_label", "Annuler"])
    ws_meta.append(["show_cancel", "true"])
    ws_meta.append(["validate_on_blur", "true"])
    ws_meta.append(["validate_on_submit", "true"])
    
    # Feuille Sections
    ws_sections = wb.create_sheet("Sections")
    ws_sections.append(["id", "title", "description", "icon", "order"])
    ws_sections.append(["contact-info", "Informations de Contact", "Vos coordonnées", "user", "1"])
    
    # Feuille Fields
    ws_fields = wb.create_sheet("Fields")
    ws_fields.append([
        "id", "name", "type", "label", "placeholder", "required",
        "section_id", "group_id", "group_title", "group_layout", "group_columns",
        "description", "tooltip", "defaultValue", "min", "max", "span"
    ])
    ws_fields.append([
        "firstName", "firstName", "text", "Prénom", "Entrez votre prénom", "true",
        "contact-info", "identity", "Identité", "grid", "2",
        "", "", "", "", "", "1"
    ])
    ws_fields.append([
        "lastName", "lastName", "text", "Nom", "Entrez votre nom", "true",
        "contact-info", "identity", "", "", "",
        "", "", "", "", "", "1"
    ])
    ws_fields.append([
        "email", "email", "email", "Email", "votre@email.com", "true",
        "contact-info", "contact", "Contact", "grid", "1",
        "", "", "", "", "", ""
    ])
    
    # Feuille Options
    ws_options = wb.create_sheet("Options")
    ws_options.append(["field_id", "label", "value", "disabled"])
    
    # Feuille Validation
    ws_validation = wb.create_sheet("Validation")
    ws_validation.append(["field_id", "type", "message", "value", "pattern"])
    ws_validation.append(["firstName", "required", "Le prénom est obligatoire", "", ""])
    ws_validation.append(["firstName", "minLength", "Minimum 2 caractères", "2", ""])
    ws_validation.append(["email", "required", "L'email est obligatoire", "", ""])
    ws_validation.append(["email", "email", "Format d'email invalide", "", ""])
    
    # Sauvegarder
    wb.save(output_path)


# Exemple d'utilisation
if __name__ == "__main__":
    # Créer un template
    create_excel_template("template_referentiel.xlsx")
    
    # Parser un fichier Excel
    # parser = ExcelReferentielParser()
    # referentiel = parser.parse_excel_file("mon_formulaire.xlsx")
    # print(json.dumps(referentiel, indent=2, ensure_ascii=False))
