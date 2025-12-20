"""
Créer un fichier Excel exemple pour tester l'import
"""
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

# Créer un nouveau workbook
wb = openpyxl.Workbook()
ws: Worksheet = wb.active  # type: ignore
ws.title = "Formulaire"

# Style pour l'en-tête
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")

# En-têtes
headers = ["Section", "Groupe", "Champ", "Type", "Label", "Requis", "Options"]
for col, header in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.fill = header_fill
    cell.font = header_font

# Données d'exemple
data = [
    ["Informations personnelles", "", "nom", "text", "Nom complet", "oui", ""],
    ["Informations personnelles", "", "prenom", "text", "Prénom", "oui", ""],
    ["Informations personnelles", "", "email", "email", "Adresse e-mail", "oui", ""],
    ["Informations personnelles", "", "telephone", "text", "Téléphone", "non", ""],
    ["Informations personnelles", "", "date_naissance", "date", "Date de naissance", "oui", ""],
    ["", "", "", "", "", "", ""],
    ["Adresse", "", "rue", "text", "Rue", "oui", ""],
    ["Adresse", "", "ville", "text", "Ville", "oui", ""],
    ["Adresse", "", "code_postal", "text", "Code postal", "oui", ""],
    ["Adresse", "", "pays", "select", "Pays", "oui", "France|Belgique|Suisse|Canada"],
    ["", "", "", "", "", "", ""],
    ["Préférences", "", "newsletter", "checkbox", "S'abonner à la newsletter", "non", ""],
    ["Préférences", "", "langue", "select", "Langue préférée", "oui", "Français|Anglais|Espagnol"],
    ["Préférences", "", "theme", "radio", "Thème de l'interface", "oui", "Clair|Sombre|Auto"],
    ["", "", "", "", "", "", ""],
    ["Commentaires", "", "message", "textarea", "Votre message", "non", ""],
    ["Commentaires", "", "note", "number", "Note (sur 10)", "non", ""],
]

# Ajouter les données
for row_idx, row_data in enumerate(data, start=2):
    for col_idx, value in enumerate(row_data, start=1):
        ws.cell(row=row_idx, column=col_idx, value=value)

# Ajuster la largeur des colonnes
ws.column_dimensions['A'].width = 30
ws.column_dimensions['B'].width = 20
ws.column_dimensions['C'].width = 20
ws.column_dimensions['D'].width = 15
ws.column_dimensions['E'].width = 30
ws.column_dimensions['F'].width = 10
ws.column_dimensions['G'].width = 40

# Sauvegarder
filename = "exemple_referentiel_simple.xlsx"
wb.save(filename)
print(f"✅ Fichier Excel créé: {filename}")
print(f"\n📋 Le fichier contient:")
print(f"   - {len([r for r in data if r[2]])} champs")
print(f"   - {len(set([r[0] for r in data if r[0]]))} sections")
print(f"\n📝 Format:")
print("   Section | Groupe | Champ | Type | Label | Requis | Options")
print(f"\n🎯 Types de champs inclus:")
types = set([r[3] for r in data if r[3]])
for t in types:
    print(f"   - {t}")
